import io
import json
import asyncio
from datetime import datetime, timezone
from typing import Optional

from fastapi import HTTPException, BackgroundTasks
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse, Response
from sqlalchemy.orm import Session
from sqlalchemy import func, or_

from app.database.session import get_db, SessionLocal
from app.models.conversation import Conversation
from app.models.customer import Customer
from app.models.message import Message
from app.models.user import User
from app.models.company_config import CompanyConfig
from app.schemas.customer_schema import (
    CustomerCreate,
    CustomerResponse,
)
from app.services.config_service import get_or_create_config
from app.services.deps import get_current_user, require_company_manager
from app.routers.config_router import _evo_config
from app.services import evolution

router = APIRouter(
    prefix="/customers",
    tags=["Customers"],
)


@router.post("/", response_model=CustomerResponse)
def create_customer(
    customer: CustomerCreate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    new_customer = Customer(
        company_id=current_user.company_id,
        name=customer.name,
        phone=customer.phone,
        email=customer.email or None,
        company=customer.company or None,
        city=customer.city or None,
        notes=customer.notes or None,
    )
    db.add(new_customer)
    db.commit()
    db.refresh(new_customer)
    return new_customer


@router.get("/")
def get_customers(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    q: Optional[str] = None,
    limit: int = 50,
    offset: int = 0,
):
    query = db.query(Customer).filter(Customer.company_id == current_user.company_id)
    if q and q.strip():
        ql = f"%{q.strip()}%"
        query = query.filter(
            or_(
                Customer.name.ilike(ql),
                Customer.phone.ilike(ql),
                Customer.email.ilike(ql),
            )
        )
    total = query.count()
    customers = query.order_by(Customer.id.desc()).offset(offset).limit(limit).all()

    customer_ids = [c.id for c in customers]
    conv_counts = (
        db.query(Conversation.customer_id, func.count(Conversation.id))
        .filter(Conversation.customer_id.in_(customer_ids))
        .group_by(Conversation.customer_id)
        .all()
    )
    count_map = dict(conv_counts)

    items = []
    for c in customers:
        items.append({
            "id": c.id,
            "company_id": c.company_id,
            "name": c.name,
            "phone": c.phone,
            "email": c.email,
            "company": c.company,
            "city": c.city,
            "conversation_count": count_map.get(c.id, 0),
        })

    return {"total": total, "items": items}


def _get_export_data(db: Session, company_id: int) -> list[dict]:
    customers = (
        db.query(Customer)
        .filter(Customer.company_id == company_id)
        .order_by(Customer.id)
        .all()
    )
    if not customers:
        return []

    customer_ids = [c.id for c in customers]
    conv_counts = (
        db.query(Conversation.customer_id, func.count(Conversation.id))
        .filter(Conversation.customer_id.in_(customer_ids))
        .group_by(Conversation.customer_id)
        .all()
    )
    count_map = dict(conv_counts)

    return [
        {
            "id": c.id,
            "name": c.name or "",
            "phone": c.phone or "",
            "email": c.email or "",
            "company": c.company or "",
            "city": c.city or "",
            "conversation_count": count_map.get(c.id, 0),
        }
        for c in customers
    ]


@router.get("/export")
def export_customers_json(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Exporta leads como JSON para download."""
    data = _get_export_data(db, current_user.company_id)
    content = json.dumps(data, ensure_ascii=False, indent=2)
    return Response(
        content=content,
        media_type="application/json",
        headers={
            "Content-Disposition": f'attachment; filename="leads_{current_user.company_id}.json"'
        },
    )


@router.get("/export/xlsx")
def export_customers_xlsx(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Exporta leads como Excel (.xlsx) para download."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise HTTPException(status_code=500, detail="Export Excel nao disponivel")

    data = _get_export_data(db, current_user.company_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4F7CFF", end_color="4F7CFF", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    headers = ["ID", "Nome", "Telefone", "Email", "Empresa", "Cidade", "Conversas"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for row_idx, item in enumerate(data, 2):
        ws.cell(row=row_idx, column=1, value=item["id"]).border = thin_border
        ws.cell(row=row_idx, column=2, value=item["name"]).border = thin_border
        ws.cell(row=row_idx, column=3, value=item["phone"]).border = thin_border
        ws.cell(row=row_idx, column=4, value=item["email"]).border = thin_border
        ws.cell(row=row_idx, column=5, value=item["company"]).border = thin_border
        ws.cell(row=row_idx, column=6, value=item["city"]).border = thin_border
        ws.cell(row=row_idx, column=7, value=item["conversation_count"]).border = thin_border

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 26
    ws.column_dimensions["E"].width = 22
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 12

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="leads_{current_user.company_id}.xlsx"'
        },
    )


def _send_bulk_background(company_id: int, phones: list[str], text: str) -> None:
    """Envia mensagem em massa pelo WhatsApp em background."""
    db: Session = SessionLocal()
    try:
        config = get_or_create_config(db, company_id)
        base_url, api_key, instance = _evo_config(config)
        if not base_url or not api_key or not instance:
            return
        for phone in phones:
            try:
                asyncio.run(evolution.send_text(
                    to_phone=phone,
                    text=text,
                    base_url=base_url,
                    api_key=api_key,
                    instance=instance,
                ))
            except Exception:
                pass
    finally:
        db.close()


@router.post("/bulk-message")
def bulk_message(
    data: dict,
    current_user: User = Depends(require_company_manager),
    db: Session = Depends(get_db),
    background_tasks: BackgroundTasks = None,
):
    """Envia mensagem em massa para leads."""
    text = (data.get("text") or "").strip()
    if not text:
        raise HTTPException(status_code=400, detail="Mensagem obrigatoria")

    customer_ids = data.get("customer_ids")
    send_all = data.get("all", False)

    if send_all:
        customers = (
            db.query(Customer)
            .filter(Customer.company_id == current_user.company_id)
            .all()
        )
    elif customer_ids:
        customers = (
            db.query(Customer)
            .filter(
                Customer.id.in_(customer_ids),
                Customer.company_id == current_user.company_id,
            )
            .all()
        )
    else:
        raise HTTPException(status_code=400, detail="Selecione os leads ou marque 'enviar para todos'")

    if not customers:
        raise HTTPException(status_code=400, detail="Nenhum lead encontrado")

    phones = [c.phone for c in customers]

    if background_tasks:
        background_tasks.add_task(_send_bulk_background, current_user.company_id, phones, text)

    return {"message": f"Mensagem enviada para {len(phones)} leads", "count": len(phones)}


@router.get("/{customer_id}", response_model=CustomerResponse)
def get_customer(
    customer_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    customer = (
        db.query(Customer)
        .filter(
            Customer.id == customer_id,
            Customer.company_id == current_user.company_id,
        )
        .first()
    )
    if not customer:
        raise HTTPException(status_code=404, detail="Customer not found")
    return customer


@router.delete("/{customer_id}")
def delete_customer(
    customer_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    customer = (
        db.query(Customer)
        .filter(
            Customer.id == customer_id,
            Customer.company_id == current_user.company_id,
        )
        .first()
    )
    if not customer:
        raise HTTPException(status_code=404, detail="Customer not found")

    conv_ids = [c.id for c in db.query(Conversation.id).filter(Conversation.customer_id == customer.id).all()]
    if conv_ids:
        db.query(Message).filter(Message.conversation_id.in_(conv_ids)).delete(synchronize_session=False)
        db.query(Conversation).filter(Conversation.id.in_(conv_ids)).delete(synchronize_session=False)

    db.delete(customer)
    db.commit()
    return {"message": "Lead removido com sucesso"}
